"""Excel workbooks (.xlsx / .xlsm): every non-empty sheet becomes a table.

Read with openpyxl (already a Bearings dependency – no DuckDB extension download, works offline).
Column types come from the cell types Excel stored, not from guessing text: a column of text cells
stays text (so `00123` keeps its leading zeros), whole numbers become BIGINT, dates DATE/TIMESTAMP,
and a column mixing types becomes VARCHAR. Formulas load as their last calculated value.

Naming: a workbook with one (selected) sheet → a table named after the file; with several →
`<file>_<sheet>`. The header is the first non-empty row unless `header_row` says otherwise.
"""
from __future__ import annotations

import csv
import datetime as dt
import os
import tempfile
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from ..db import safe_name
from .base import FileReader, LoadItem, LoadOptions, sql_str

SCAN_ROWS = 1000  # rows looked at to decide whether a sheet is empty


def _open(p: Path):
    import openpyxl
    return openpyxl.load_workbook(p, read_only=True, data_only=True)


def _has_data(ws) -> bool:
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if any(v is not None and str(v).strip() != "" for v in row):
            return True
        if i >= SCAN_ROWS:
            return False
    return False


def _kind(v) -> str:
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, int):
        return "int" if -(2 ** 63) <= v < 2 ** 63 else "float"
    if isinstance(v, float):
        return "float"
    if isinstance(v, dt.datetime):
        return "date" if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0) else "timestamp"
    if isinstance(v, dt.date):
        return "date"
    if isinstance(v, dt.time):
        return "time"
    return "str"


def column_type(kinds: set[str]) -> str:
    """DuckDB type for a column from the set of cell kinds seen in it."""
    k = kinds - {"none"}
    if not k or "str" in k:
        return "VARCHAR"
    if k == {"bool"}:
        return "BOOLEAN"
    if k == {"int"}:
        return "BIGINT"
    if k <= {"int", "float"}:
        return "DOUBLE"
    if k == {"date"}:
        return "DATE"
    if k <= {"date", "timestamp"}:
        return "TIMESTAMP"
    if k == {"time"}:
        return "TIME"
    return "VARCHAR"


def as_text(v) -> str | None:
    """How a cell is written to the staging CSV (and what a VARCHAR column shows)."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if _kind(v) == "date" else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    return str(v)


def header_names(raw: list) -> list[str]:
    """Header cells → unique, non-empty column names (original spelling kept, like CSV headers)."""
    out, used = [], set()
    for i, h in enumerate(raw):
        base = " ".join(str(as_text(h)).split()) if h is not None and str(h).strip() else f"column_{i + 1}"
        name, n = base, 1
        while name.lower() in used:
            n += 1
            name = f"{base}_{n}"
        used.add(name.lower())
        out.append(name)
    return out


def read_sheet(ws, header_row: int | None = None) -> tuple[list[str], list[list], list[set[str]]]:
    """(column names, rows, cell kinds per column) of one worksheet. Empty rows are skipped; empty trailing
    columns are dropped."""
    header, rows = None, []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [None if (v is None or (isinstance(v, str) and v.strip() == "")) else v for v in row]
        if header is None:
            if (header_row and i == header_row) or (not header_row and any(v is not None for v in vals)):
                header = list(vals)
            continue
        if any(v is not None for v in vals):
            rows.append(vals)
    if header is None:
        return [], [], []
    width = max([i + 1 for i, h in enumerate(header) if h is not None] +
                [max((i + 1 for i, v in enumerate(r) if v is not None), default=0) for r in rows] + [0])
    names = header_names((header + [None] * width)[:width])
    kinds: list[set[str]] = [set() for _ in range(width)]
    for r in rows:
        r[:] = (r + [None] * width)[:width]
        for j, v in enumerate(r):
            if v is not None:
                kinds[j].add(_kind(v))
    return names, rows, kinds


class ExcelReader(FileReader):
    name = "excel"
    label = "Excel workbook"
    extensions = (".xlsx", ".xlsm", ".xls")
    needs = "openpyxl"

    def items_for_file(self, p: Path, opts: LoadOptions) -> list[LoadItem]:
        if p.name.lower().endswith(".xls"):
            raise ValueError(f"{p.name}: legacy .xls isn't supported – open it in Excel and save as .xlsx")
        wb = _open(p)
        try:
            titles = [ws.title for ws in wb.worksheets]
            want = {s.lower() for s in opts.sheets}
            missing = want - {t.lower() for t in titles}
            if want and missing == want:
                return []  # none of the requested sheets is in this workbook (another file may have them)
            chosen = [t for t in titles if (not want or t.lower() in want) and _has_data(wb[t])]
        finally:
            wb.close()
        stem = self.stem(p)
        return [LoadItem(safe_name(stem if len(chosen) == 1 else f"{stem}_{t}"), f"{p}#{t}", self.name, [p], {"sheet": t})
                for t in chosen]

    @contextmanager
    def relation(self, con, item: LoadItem, opts: LoadOptions) -> Iterator[str]:
        wb = _open(item.paths[0])
        try:
            names, rows, kinds = read_sheet(wb[item.options["sheet"]], opts.header_row)
        finally:
            wb.close()
        if not names:
            raise ValueError(f"sheet '{item.options['sheet']}' is empty")
        types = ["VARCHAR" if opts.all_varchar else column_type(k) for k in kinds]
        fd, tmp = tempfile.mkstemp(prefix="bearings-xlsx-", suffix=".csv")
        try:
            with os.fdopen(fd, "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(names)
                for r in rows:
                    w.writerow(["" if v is None else as_text(v) for v in r])
            cols = "{" + ", ".join(f"{sql_str(n)}: {sql_str(t)}" for n, t in zip(names, types)) + "}"
            yield (f"read_csv({sql_str(tmp)}, header=true, auto_detect=false, delim=',', quote='\"', escape='\"', "
                   f"columns={cols}, nullstr='')")
        finally:
            try:
                os.unlink(tmp)
            except OSError:  # pragma: no cover
                pass
